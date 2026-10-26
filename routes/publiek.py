from datetime import date, datetime
from io import BytesIO
import math
from urllib.parse import quote_plus

from fastapi import APIRouter, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import bereken_stortingsbedrag, get_db
from auth import (
    configured_password_hash,
    configured_username,
    verify_password,
)


router = APIRouter()
templates = Jinja2Templates(directory="templates")


# ============================================================
# EURO
# ============================================================

def euro(value):

    return (
        f"€ {float(value or 0):,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


templates.env.filters["euro"] = euro


# ============================================================
# TOTALEN
# ============================================================

def totals(row, nummer):

    bonnen = (
        (row[f"c{nummer}_blauw"] or 0) * 0.75
        +
        (row[f"c{nummer}_groen"] or 0) * 1.00
    )

    briefgeld = (
        (row[f"c{nummer}_munt_5"] or 0) * 5
        +
        (row[f"c{nummer}_munt_10"] or 0) * 10
        +
        (row[f"c{nummer}_munt_20"] or 0) * 20
        +
        (row[f"c{nummer}_munt_50"] or 0) * 50
    )

    muntgeld = (
        row[f"c{nummer}_overig"] or 0
    )

    contant = (
        briefgeld
        +
        muntgeld
    )

    totaal = (
        bonnen
        +
        contant
    )

    return (
        bonnen,
        muntgeld,
        briefgeld,
        contant,
        totaal,
    )


def derde_collecte_actief(row):

    return bool(
        str(row["doel_3"] or "").strip()
    )


def vierde_collecte_actief(row):

    return bool(
        str(row["doel_4"] or "").strip()
    )


COLLECTE_VELDEN = (
    "blauw",
    "groen",
    "munt_5",
    "munt_10",
    "munt_20",
    "munt_50",
    "overig",
)


# ============================================================
# LOGIN
# ============================================================

@router.get(
    "/login",
    response_class=HTMLResponse
)
def login_page(
    request: Request,
    fout: str | None = None,
):

    if request.session.get("logged_in"):

        return RedirectResponse(
            "/",
            status_code=303,
        )

    return templates.TemplateResponse(
        request=request,
        name="login.html",
        context={
            "fout": fout,
        },
    )


@router.post("/login")
def login(
    request: Request,
    gebruikersnaam: str = Form(...),
    wachtwoord: str = Form(...),
):

    password_hash = (
        configured_password_hash()
    )

    geldig = (
        bool(password_hash)
        and gebruikersnaam.strip()
        == configured_username()
        and verify_password(
            wachtwoord,
            password_hash,
        )
    )

    if not geldig:

        return templates.TemplateResponse(
            request=request,
            name="login.html",
            context={
                "fout": (
                    "Gebruikersnaam of wachtwoord "
                    "is niet juist."
                ),
            },
            status_code=401,
        )

    request.session.clear()

    request.session["logged_in"] = True
    request.session["username"] = (
        gebruikersnaam.strip()
    )

    return RedirectResponse(
        "/",
        status_code=303,
    )


@router.post("/logout")
def logout(request: Request):

    request.session.clear()

    return RedirectResponse(
        "/login",
        status_code=303,
    )


# ============================================================
# HOME
# ============================================================

@router.get(
    "/",
    response_class=HTMLResponse
)
def home(request: Request):

    vandaag = date.today().isoformat()

    with get_db() as db:

        doel = db.execute(
            """
            SELECT *
            FROM collectedoelen
            WHERE datum=?
            """,
            (vandaag,),
        ).fetchone()

        recent = db.execute(
            """
            SELECT *
            FROM tellingen
            ORDER BY datum DESC
            LIMIT 8
            """
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "vandaag": vandaag,
            "doel": doel,
            "recent": recent,
        },
    )


# ============================================================
# TELLEN
# ============================================================

@router.get(
    "/tellen",
    response_class=HTMLResponse
)
def tellen(
    request: Request,
    datum: str | None = None,
):

    datum = (
        datum
        or date.today().isoformat()
    )

    with get_db() as db:

        doel = db.execute(
            """
            SELECT *
            FROM collectedoelen
            WHERE datum=?
            """,
            (datum,),
        ).fetchone()

        bestaand = db.execute(
            """
            SELECT *
            FROM tellingen
            WHERE datum=?
            """,
            (datum,),
        ).fetchone()

        rows_bevestigingen = db.execute(
            """
            SELECT *
            FROM collecte_bevestigingen
            WHERE datum=?
            """,
            (datum,),
        ).fetchall()

    collecties = {}
    bevestigingen = {}

    for nummer in range(1, 5):
        bevestiging = next(
            (
                row
                for row in rows_bevestigingen
                if row["collecte"] == nummer
            ),
            None,
        )
        bevestigingen[nummer] = bevestiging is not None

        if bevestiging:
            collecties[nummer] = {
                "doelcode": bevestiging["doelcode"],
                "doel": bevestiging["doel"],
                **{
                    veld: bevestiging[veld]
                    for veld in COLLECTE_VELDEN
                },
            }
        else:
            collecties[nummer] = {
                "doelcode": (
                    bestaand[f"doelcode_{nummer}"]
                    if bestaand
                    else (
                        doel[f"doelcode_{nummer}"]
                        if doel
                        else ""
                    )
                ),
                "doel": (
                    bestaand[f"doel_{nummer}"]
                    if bestaand
                    else (
                        doel[f"doel_{nummer}"]
                        if doel
                        else ""
                    )
                ),
                **{
                    veld: (
                        bestaand[f"c{nummer}_{veld}"]
                        if bestaand
                        else 0
                    )
                    for veld in COLLECTE_VELDEN
                },
            }

    return templates.TemplateResponse(
        request=request,
        name="tellen.html",
        context={
            "datum": datum,
            "doel": doel,
            "bestaand": bestaand,
            "collecties": collecties,
            "bevestigingen": bevestigingen,
        },
    )


@router.post("/tellen")
def telling_opslaan(

    datum: str = Form(...),

    doelcode_1: str = Form(""),
    doel_1: str = Form(...),

    doelcode_2: str = Form(""),
    doel_2: str = Form(...),

    doelcode_3: str = Form(""),
    doel_3: str = Form(""),
    doelcode_4: str = Form(""),
    doel_4: str = Form(""),

    c1_blauw: int = Form(0),
    c1_groen: int = Form(0),
    c1_munt_5: int = Form(0),
    c1_munt_10: int = Form(0),
    c1_munt_20: int = Form(0),
    c1_munt_50: int = Form(0),
    c1_overig: float = Form(0),

    c2_blauw: int = Form(0),
    c2_groen: int = Form(0),
    c2_munt_5: int = Form(0),
    c2_munt_10: int = Form(0),
    c2_munt_20: int = Form(0),
    c2_munt_50: int = Form(0),
    c2_overig: float = Form(0),

    c3_blauw: int = Form(0),
    c3_groen: int = Form(0),
    c3_munt_5: int = Form(0),
    c3_munt_10: int = Form(0),
    c3_munt_20: int = Form(0),
    c3_munt_50: int = Form(0),
    c3_overig: float = Form(0),

    c4_blauw: int = Form(0),
    c4_groen: int = Form(0),
    c4_munt_5: int = Form(0),
    c4_munt_10: int = Form(0),
    c4_munt_20: int = Form(0),
    c4_munt_50: int = Form(0),
    c4_overig: float = Form(0),

    extra_storting: float = Form(0),
    opmerkingen: str = Form(""),
    geteld_door: str = Form(...),
):

    data = locals().copy()

    with get_db() as db:
        db.execute("BEGIN IMMEDIATE")
        bevestigingen = db.execute(
            """
            SELECT *
            FROM collecte_bevestigingen
            WHERE datum=?
            """,
            (datum,),
        ).fetchall()

        for bevestiging in bevestigingen:
            nummer = bevestiging["collecte"]
            beschermde_velden = {
                f"doelcode_{nummer}": bevestiging["doelcode"],
                f"doel_{nummer}": bevestiging["doel"],
                **{
                    f"c{nummer}_{veld}": bevestiging[veld]
                    for veld in COLLECTE_VELDEN
                },
            }

            for veld, bewaarde_waarde in beschermde_velden.items():
                waarde = data[veld]
                if isinstance(bewaarde_waarde, str):
                    komt_overeen = waarde.strip() == bewaarde_waarde
                else:
                    komt_overeen = waarde == bewaarde_waarde

                if not komt_overeen:
                    return HTMLResponse(
                        "Deze collecte is bevestigd en vergrendeld. "
                        "Kies eerst 'Telling wijzigen'.",
                        status_code=409,
                    )

        columns = list(data.keys())

        placeholders = ",".join(
            "?"
            for _ in columns
        )

        updates = ",".join(
            f"{column}=excluded.{column}"
            for column in columns
            if column != "datum"
        )

        db.execute(
            f"""
            INSERT INTO tellingen (
                {','.join(columns)}
            )
            VALUES ({placeholders})

            ON CONFLICT(datum)
            DO UPDATE SET
                {updates}
            """,
            tuple(
                data[column]
                for column in columns
            ),
        )

        db.commit()

    return RedirectResponse(
        f"/controle/{datum}",
        status_code=303,
    )


@router.post("/tellen/bevestigen/{nummer}")
async def collecte_bevestigen(
    nummer: int,
    request: Request,
):

    if nummer not in range(1, 5):
        return JSONResponse(
            {"error": "Ongeldig collectenummer."},
            status_code=404,
        )

    form = await request.form()
    datum = str(form.get("datum", "")).strip()

    try:
        date.fromisoformat(datum)
        waarden = {
            veld: (
                float(str(form.get(f"c{nummer}_{veld}", "0") or "0"))
                if veld == "overig"
                else int(str(form.get(f"c{nummer}_{veld}", "0") or "0"))
            )
            for veld in COLLECTE_VELDEN
        }
    except (TypeError, ValueError):
        return JSONResponse(
            {"error": "Datum of aantallen zijn ongeldig."},
            status_code=422,
        )

    if any(waarde < 0 for waarde in waarden.values()):
        return JSONResponse(
            {"error": "Aantallen en bedragen mogen niet negatief zijn."},
            status_code=422,
        )

    if not math.isfinite(waarden["overig"]):
        return JSONResponse(
            {"error": "Muntgeld moet een geldig bedrag zijn."},
            status_code=422,
        )

    doelcode = str(form.get(f"doelcode_{nummer}", "")).strip()
    doel = str(form.get(f"doel_{nummer}", "")).strip()

    if nummer < 3 and not doel:
        return JSONResponse(
            {"error": "Vul eerst de omschrijving van deze collecte in."},
            status_code=422,
        )

    with get_db() as db:
        db.execute("BEGIN IMMEDIATE")
        bestaand = db.execute(
            """
            SELECT 1
            FROM collecte_bevestigingen
            WHERE datum=? AND collecte=?
            """,
            (datum, nummer),
        ).fetchone()

        if bestaand:
            return JSONResponse(
                {"error": "Deze collecte is al bevestigd."},
                status_code=409,
            )

        db.execute(
            """
            INSERT INTO collecte_bevestigingen (
                datum, collecte, doelcode, doel,
                blauw, groen, munt_5, munt_10, munt_20, munt_50, overig
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datum,
                nummer,
                doelcode,
                doel,
                waarden["blauw"],
                waarden["groen"],
                waarden["munt_5"],
                waarden["munt_10"],
                waarden["munt_20"],
                waarden["munt_50"],
                waarden["overig"],
            ),
        )
        db.commit()

    return JSONResponse({"bevestigd": True})


@router.post("/tellen/wijzigen/{nummer}")
async def collecte_wijzigen(
    nummer: int,
    request: Request,
):

    if nummer not in range(1, 5):
        return JSONResponse(
            {"error": "Ongeldig collectenummer."},
            status_code=404,
        )

    form = await request.form()
    datum = str(form.get("datum", "")).strip()

    try:
        date.fromisoformat(datum)
    except ValueError:
        return JSONResponse(
            {"error": "Datum is ongeldig."},
            status_code=422,
        )

    with get_db() as db:
        resultaat = db.execute(
            """
            DELETE FROM collecte_bevestigingen
            WHERE datum=? AND collecte=?
            """,
            (datum, nummer),
        )

        if resultaat.rowcount == 0:
            return JSONResponse(
                {"error": "Deze collecte is niet bevestigd."},
                status_code=409,
            )

        db.commit()

    return JSONResponse({"bevestigd": False})


# ============================================================
# CONTROLE
# ============================================================

@router.get(
    "/controle/{datum}",
    response_class=HTMLResponse
)
def controle(
    request: Request,
    datum: str,
):

    with get_db() as db:

        row = db.execute(
            """
            SELECT *
            FROM tellingen
            WHERE datum=?
            """,
            (datum,),
        ).fetchone()

    if not row:

        return RedirectResponse(
            "/",
            status_code=303,
        )

    t1 = totals(row, 1)
    t2 = totals(row, 2)
    t3 = totals(row, 3)
    t4 = totals(row, 4)

    heeft_derde = (
        derde_collecte_actief(row)
    )
    heeft_vierde = (
        vierde_collecte_actief(row)
    )

    totaal = (
        t1[4]
        +
        t2[4]
        +
        (
            t3[4]
            if heeft_derde
            else 0
        )
        +
        (
            t4[4]
            if heeft_vierde
            else 0
        )
        +
        (row["extra_storting"] or 0)
    )
    stortingsbedrag = bereken_stortingsbedrag(row)

    return templates.TemplateResponse(
        request=request,
        name="controle.html",
        context={
            "r": row,
            "t1": t1,
            "t2": t2,
            "t3": t3,
            "t4": t4,
            "heeft_derde": heeft_derde,
            "heeft_vierde": heeft_vierde,
            "totaal": totaal,
            "stortingsbedrag": stortingsbedrag,
        },
    )


# ============================================================
# COLLECTEDOELEN
# ============================================================

@router.get(
    "/doelen",
    response_class=HTMLResponse
)
def doelen(
    request: Request,
    melding: str | None = None,
    fout: str | None = None,
):

    with get_db() as db:

        rows = db.execute(
            """
            SELECT *
            FROM collectedoelen
            ORDER BY datum ASC
            """
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="doelen.html",
        context={
            "rows": rows,
            "melding": melding,
            "fout": fout,
        },
    )


# ============================================================
# COLLECTEDOEL HANDMATIG OPSLAAN
# ============================================================

@router.post("/doelen")
def doel_opslaan(

    datum: str = Form(...),

    doel_1: str = Form(...),

    doel_2: str = Form(...),

    doel_3: str = Form(""),
    doel_4: str = Form(""),
):

    with get_db() as db:

        db.execute(
            """
            INSERT INTO collectedoelen(
                datum,
                doel_1,
                doel_2,
                doel_3,
                doel_4
            )
            VALUES(?,?,?,?,?)

            ON CONFLICT(datum)
            DO UPDATE SET
                doel_1=excluded.doel_1,
                doel_2=excluded.doel_2,
                doel_3=excluded.doel_3,
                doel_4=excluded.doel_4
            """,
            (
                datum,
                doel_1.strip(),
                doel_2.strip(),
                doel_3.strip(),
                doel_4.strip(),
            ),
        )

        db.commit()

    return RedirectResponse(
        "/doelen",
        status_code=303,
    )


# ============================================================
# DATUM UIT EXCEL
# ============================================================

def excel_datum_naar_iso(
    raw_datum,
    rij,
):

    if raw_datum is None:

        raise ValueError(
            f"Datum ontbreekt op rij {rij}."
        )

    if isinstance(
        raw_datum,
        datetime
    ):

        return (
            raw_datum
            .date()
            .isoformat()
        )

    if isinstance(
        raw_datum,
        date
    ):

        return raw_datum.isoformat()

    tekst = str(
        raw_datum
    ).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
    ):

        try:

            return (
                datetime.strptime(
                    tekst,
                    fmt,
                )
                .date()
                .isoformat()
            )

        except ValueError:

            pass

    raise ValueError(
        f"Ongeldige datum op rij {rij}: {tekst}"
    )


# ============================================================
# EXCELWAARDE NAAR TEKST
# ============================================================

def excel_tekst(value):

    if value is None:

        return ""

    # Voorkomt bijvoorbeeld 101.0
    # wanneer Excel een geheel getal als float aanlevert.

    if (
        isinstance(value, float)
        and value.is_integer()
    ):

        return str(int(value))

    return str(value).strip()


# ============================================================
# COLLECTEDOELEN IMPORTEREN
# ============================================================

@router.post("/doelen/importeren")
async def doelen_importeren(
    bestand: UploadFile = File(...),
):

    if (
        not bestand.filename
        or not bestand.filename
        .lower()
        .endswith(".xlsx")
    ):

        return RedirectResponse(
            (
                "/doelen?"
                "fout=Gebruik+een+"
                "Excelbestand+(.xlsx)"
            ),
            status_code=303,
        )

    try:

        inhoud = await bestand.read()

        wb = load_workbook(
            BytesIO(inhoud),
            data_only=True,
        )

        ws = wb.active

        werkelijke_headers = []

        for kolom in range(1, ws.max_column + 1):

            waarde = ws.cell(
                1,
                kolom,
            ).value

            werkelijke_headers.append(
                str(
                    waarde or ""
                )
                .strip()
                .lower()
            )

        while werkelijke_headers and not werkelijke_headers[-1]:
            werkelijke_headers.pop()

        omschrijving_headers = {
            tuple(["datum", "collecte 1", "collecte 2"]): (2, 3),
            tuple(["datum", "collecte 1", "collecte 2", "collecte 3"]): (2, 3, 4),
            tuple(["datum", "collecte 1", "collecte 2", "collecte 3", "collecte 4"]): (
                2, 3, 4, 5
            ),
        }
        legacy_headers = {
            tuple([
                "datum", "doel", "collecte 1", "doel", "collecte 2",
                "doel", "collecte 3",
            ]): ((2, 4, 6), (3, 5, 7)),
            tuple([
                "datum", "doel", "collecte 1", "doel", "collecte 2",
                "doel", "collecte 3", "doel", "collecte 4",
            ]): ((2, 4, 6, 8), (3, 5, 7, 9)),
        }

        header = tuple(werkelijke_headers)
        if header in omschrijving_headers:
            doel_kolommen = None
            omschrijving_kolommen = omschrijving_headers[header]
        elif header in legacy_headers:
            doel_kolommen, omschrijving_kolommen = legacy_headers[header]
        else:

            raise ValueError(
                "De eerste rij moet bevatten: Datum | Collecte 1 | "
                "Collecte 2, eventueel aangevuld met Collecte 3 en Collecte 4. "
                "Bestaande bestanden met Doel-codekolommen blijven ook geldig."
            )

        aantal_collectes = len(omschrijving_kolommen)

        aantal = 0

        with get_db() as db:

            for rij in range(
                2,
                ws.max_row + 1,
            ):

                raw_datum = (
                    ws.cell(rij, 1).value
                )

                omschrijvingen = [
                    excel_tekst(ws.cell(rij, kolom).value)
                    for kolom in omschrijving_kolommen
                ]
                doelcodes = (
                    [
                        excel_tekst(ws.cell(rij, kolom).value)
                        for kolom in doel_kolommen
                    ]
                    if doel_kolommen
                    else [""] * aantal_collectes
                )
                doelcodes = [
                    code if omschrijving else ""
                    for code, omschrijving in zip(doelcodes, omschrijvingen)
                ]

                # Helemaal lege regel overslaan

                if (
                    raw_datum is None
                    and not any(doelcodes)
                    and not any(omschrijvingen)
                ):

                    continue

                datum = excel_datum_naar_iso(
                    raw_datum,
                    rij,
                )

                if not omschrijvingen[0]:

                    raise ValueError(
                        f"Collecte 1 ontbreekt "
                        f"op rij {rij}."
                    )

                if len(omschrijvingen) < 2 or not omschrijvingen[1]:

                    raise ValueError(
                        f"Collecte 2 ontbreekt "
                        f"op rij {rij}."
                    )

                kolommen = [
                    "datum",
                    "doel_1",
                    "doel_2",
                    "doel_3",
                ]
                waarden = [
                    datum,
                    omschrijvingen[0],
                    omschrijvingen[1],
                    omschrijvingen[2] if len(omschrijvingen) > 2 else "",
                ]

                if aantal_collectes == 4:
                    kolommen.append("doel_4")
                    waarden.append(omschrijvingen[3])

                if doel_kolommen:
                    code_kolommen = [
                        f"doelcode_{nummer}"
                        for nummer in range(1, aantal_collectes + 1)
                    ]
                    kolommen.extend(code_kolommen)
                    waarden.extend(doelcodes)

                updates = ",".join(
                    (
                        f"{kolom}=CASE WHEN excluded.{kolom}='' "
                        f"THEN collectedoelen.{kolom} "
                        f"ELSE excluded.{kolom} END"
                        if kolom.startswith("doelcode_")
                        else f"{kolom}=excluded.{kolom}"
                    )
                    for kolom in kolommen
                    if kolom != "datum"
                )

                db.execute(
                    f"""
                    INSERT INTO collectedoelen({','.join(kolommen)})
                    VALUES({','.join('?' for _ in kolommen)})
                    ON CONFLICT(datum) DO UPDATE SET {updates}
                    """,
                    waarden,
                )

                aantal += 1

            db.commit()

        return RedirectResponse(
            (
                "/doelen?"
                f"melding={aantal}+"
                "zondagen+geimporteerd"
            ),
            status_code=303,
        )

    except Exception as exc:

        return RedirectResponse(
            (
                "/doelen?fout="
                f"{quote_plus(str(exc))}"
            ),
            status_code=303,
        )


# ============================================================
# VOORBEELD EXCEL COLLECTEDOELEN
# ============================================================

@router.get("/doelen/voorbeeld")
def doelen_voorbeeld():

    wb = Workbook()

    ws = wb.active
    ws.title = "Collectedoelen"

    ws.append(
        [
            "Datum",
            "Collecte 1",
            "Collecte 2",
            "Collecte 3",
            "Collecte 4",
        ]
    )

    jaar = date.today().year

    ws.append(
        [
            date(jaar, 1, 4),
            "Kerk",
            "Diaconie",
            "",
            "",
        ]
    )

    ws.append(
        [
            date(jaar, 1, 11),
            "Kerk",
            "Diaconie",
            "Extra doel",
            "Vierde doel",
        ]
    )

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="166534",
        )

    for column, width in {
        "A": 16,
        "B": 34,
        "C": 34,
        "D": 34,
        "E": 34,
    }.items():
        ws.column_dimensions[column].width = width

    for cell in ws["A"][1:]:

        cell.number_format = (
            "dd-mm-yyyy"
        )

    bio = BytesIO()

    wb.save(bio)
    bio.seek(0)

    return StreamingResponse(
        bio,
        media_type=(
            "application/vnd."
            "openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename='
                '"Voorbeeld_collectedoelen.xlsx"'
            )
        },
    )


# ============================================================
# EXCEL MAKEN
# ============================================================

def make_excel(rows):

    wb = Workbook()

    ws = wb.active
    ws.title = "Collectes"

    headers = [
        "Datum",
        "Collecte",
        "Doel",
        "Omschrijving",
        "Bonnen blauw (€0,75)",
        "Bonnen groen (€1,00)",
        "Totaal bonnen",
        "Totaal muntgeld",
        "Totaal briefgeld",
        "Totaal contant",
        "Totaal collecte",
        "Extra storting",
        "Stortingsbedrag",
        "Geteld door",
        "Opmerkingen",
    ]

    ws.append(headers)

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="166534",
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in rows:

        nummers = [1, 2]

        if derde_collecte_actief(row):

            nummers.append(3)

        if vierde_collecte_actief(row):

            nummers.append(4)

        for nummer in nummers:

            (
                bonnen,
                muntgeld,
                briefgeld,
                contant,
                totaal,
            ) = totals(
                row,
                nummer,
            )

            ws.append(
                [
                    row["datum"],

                    nummer,

                    row[
                        f"doelcode_{nummer}"
                    ],

                    row[
                        f"doel_{nummer}"
                    ],

                    row[
                        f"c{nummer}_blauw"
                    ],

                    row[
                        f"c{nummer}_groen"
                    ],

                    bonnen,
                    muntgeld,
                    briefgeld,
                    contant,
                    totaal,

                    (
                        row["extra_storting"]
                        if nummer == 1
                        else 0
                    ),

                    bereken_stortingsbedrag(row, nummer),

                    row["geteld_door"],

                    row["opmerkingen"],
                ]
            )

    # Geldkolommen G t/m M

    for excel_row in ws.iter_rows(
        min_row=2
    ):

        for index in (
            7,
            8,
            9,
            10,
            11,
            12,
            13,
        ):

            excel_row[
                index - 1
            ].number_format = (
                "€ #,##0.00"
            )

    widths = [
        14,
        10,
        14,
        32,
        20,
        20,
        17,
        17,
        18,
        18,
        18,
        17,
        19,
        24,
        42,
    ]

    for index, width in enumerate(
        widths,
        1,
    ):

        ws.column_dimensions[
            get_column_letter(index)
        ].width = width

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )

    bio = BytesIO()

    wb.save(bio)
    bio.seek(0)

    return bio


# ============================================================
# EXCEL PER ZONDAG
# ============================================================

@router.get("/excel/{datum}")
def excel_dag(datum: str):

    with get_db() as db:

        rows = db.execute(
            """
            SELECT *
            FROM tellingen
            WHERE datum=?
            """,
            (datum,),
        ).fetchall()

    bio = make_excel(rows)

    return StreamingResponse(
        bio,
        media_type=(
            "application/vnd."
            "openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                "attachment; "
                f'filename="Collectetelling_{datum}.xlsx"'
            )
        },
    )


# ============================================================
# EXCEL JAAROVERZICHT
# ============================================================

@router.get("/excel")
def excel_jaar(
    jaar: int | None = None,
):

    jaar = (
        jaar
        or date.today().year
    )

    with get_db() as db:

        rows = db.execute(
            """
            SELECT *
            FROM tellingen
            WHERE datum LIKE ?
            ORDER BY datum
            """,
            (f"{jaar}%",),
        ).fetchall()

    bio = make_excel(rows)

    return StreamingResponse(
        bio,
        media_type=(
            "application/vnd."
            "openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                "attachment; "
                f'filename="Collectes_{jaar}.xlsx"'
            )
        },
    )